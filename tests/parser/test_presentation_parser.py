from __future__ import annotations

import shutil
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation

from lumberjack import Lumberjack
from lumberjack.parser import AutoParser, OdpParser, PptParser, PptxParser


def _pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Overview"
    slide.shapes.add_textbox(0, 0, 1000000, 100000).text = "First object"
    slide.shapes.add_textbox(0, 150000, 1000000, 100000).text = "Second object"
    slide.notes_slide.notes_text_frame.text = "Speaker notes"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _odp() -> bytes:
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.presentation")
        archive.writestr(
            "content.xml",
            '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
            'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
            "<office:body><office:presentation>"
            '<draw:page draw:name="page1"><draw:frame presentation:class="title">'
            "<text:p>ODP Title</text:p></draw:frame>"
            "<draw:frame><text:p>First</text:p></draw:frame>"
            "<draw:frame><text:p>Second</text:p></draw:frame>"
            "<presentation:notes><text:p>Speaker text</text:p></presentation:notes>"
            "</draw:page></office:presentation></office:body></office:document-content>",
        )
    return output.getvalue()


def test_pptx_preserves_slide_object_and_note_order() -> None:
    data = _pptx()
    tree = PptxParser().parse(data, source_path="slides.pptx")
    result = Lumberjack().saw(data, format="pptx", source_path="slides.pptx")

    assert tree.root.children[0].title == "Overview"
    blocks = tree.root.children[0].blocks
    assert [block.text for block in blocks] == [
        "First object",
        "Second object",
        "Speaker notes",
    ]
    assert blocks[-1].attrs["role"] == "speaker_notes"
    assert [block.source_locations[0].page_start for block in blocks] == [1, 1, 1]
    assert any(
        location.page_start == 1
        for chunk in result.chunks
        for location in chunk.source_locations
    )
    assert AutoParser().parse(data).root.children[0].title == "Overview"


def test_odp_preserves_slide_object_and_note_order() -> None:
    data = _odp()
    tree = OdpParser().parse(data, source_path="slides.odp")

    assert tree.root.children[0].title == "ODP Title"
    assert [block.text for block in tree.root.children[0].blocks] == [
        "First",
        "Second",
        "Speaker text",
    ]
    assert tree.root.children[0].blocks[-1].attrs["role"] == "speaker_notes"
    assert [
        block.source_locations[0].page_start for block in tree.root.children[0].blocks
    ] == [1, 1, 1]
    assert AutoParser().parse(data).root.children[0].title == "ODP Title"


def test_ppt_requires_local_libreoffice(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "lumberjack.parser.presentation.parser.shutil.which", lambda _: None
    )
    with pytest.raises(RuntimeError, match="LibreOffice"):
        PptParser().parse(b"legacy content")


def test_ppt_uses_converted_pptx(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "lumberjack.parser.presentation.parser.shutil.which",
        lambda _: "/usr/bin/soffice",
    )

    def convert(args: list[str], **_kwargs: object) -> None:
        Path(args[-2], "input.pptx").write_bytes(_pptx())

    monkeypatch.setattr("lumberjack.parser.presentation.parser.subprocess.run", convert)
    tree = PptParser().parse(b"legacy content", source_path="slides.ppt")

    assert tree.root.children[0].title == "Overview"
    assert tree.root.children[0].blocks[0].source_locations[0].source == "slides.ppt"


def test_public_ppt_fixture_with_libreoffice() -> None:
    if not (shutil.which("libreoffice") or shutil.which("soffice")):
        pytest.skip("LibreOffice is unavailable")
    fixture = Path(__file__).parents[1] / "corpora" / "basic_test_ppt_file.ppt"
    tree = PptParser().parse(fixture.read_bytes(), source_path=fixture.name)
    assert tree.root.children
    assert tree.root.children[0].source_locations[0].page_start == 1
