"""Presentation parsers with slide and object provenance."""

from __future__ import annotations

import shutil
import subprocess
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any
from zipfile import BadZipFile, ZipFile

from ..._internal.xml_safe import parse_untrusted_xml
from ...models import DocTree, Document, InputFormat, SourceLocation
from ..builder import DocTreeBuilder


def _document(
    source: Document | bytes,
    format: InputFormat,
    title: str | None,
    metadata: dict[str, object] | None,
    path: str | Path | None,
) -> Document:
    if isinstance(source, Document):
        return source
    return Document(
        source,
        format=format,
        document_title=title,
        metadata_overrides=dict(metadata or {}),
        source_path=path,
    )


def _location(document: Document, page: int, element_id: str) -> SourceLocation:
    return SourceLocation(
        source=str(document.source_path) if document.source_path is not None else None,
        page_start=page,
        page_end=page,
        element_id=element_id,
    )


def _table_text(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    escaped = [
        [cell.replace("|", "\\|").replace("\n", " ") for cell in row]
        + [""] * (width - len(row))
        for row in rows
    ]
    lines = [
        "| " + " | ".join(escaped[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in escaped[1:])
    return "\n".join(lines)


class PptxParser:
    """Read PPTX slide shapes in XML order with notes after slide content."""

    block_kinds = frozenset({"paragraph", "table"})

    def parse(
        self,
        document: Document | bytes,
        *,
        document_title: str | None = None,
        metadata_overrides: dict[str, object] | None = None,
        source_path: str | Path | None = None,
    ) -> DocTree:
        document = _document(
            document, "pptx", document_title, metadata_overrides, source_path
        )
        if not isinstance(document.source, bytes | bytearray):
            raise TypeError("PptxParser.parse expects Document[bytes]")
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError(
                "PPTX parsing requires the 'presentations' extra: "
                "pip install 'lumberjack-py[presentations]'"
            ) from exc
        try:
            presentation = Presentation(BytesIO(bytes(document.source)))
            builder = DocTreeBuilder(
                title=document.document_title or "Anonymous",
                source_path=document.source_path,
                metadata=document.metadata_overrides,
            )
            source_parts: list[str] = []
            for number, slide in enumerate(presentation.slides, 1):
                title_shape = slide.shapes.title
                title = (
                    (title_shape.text or "").strip() if title_shape is not None else ""
                )
                builder.add_section(
                    1,
                    title or f"Slide {number}",
                    location=_location(document, number, f"slide:{number}"),
                )
                source_parts.append(title or f"Slide {number}")
                for shape in slide.shapes:
                    self._add_shape(
                        builder, shape, document, number, title_shape, source_parts
                    )
                if slide.has_notes_slide:
                    frame = slide.notes_slide.notes_text_frame
                    notes = (frame.text or "").strip() if frame is not None else ""
                    if notes:
                        builder.add_block(
                            "paragraph",
                            notes,
                            locations=(
                                _location(document, number, f"slide:{number}:notes"),
                            ),
                            attrs={"role": "speaker_notes"},
                        )
                        source_parts.append(notes)
            return replace(builder.build(), source="\n\n".join(source_parts))
        except Exception as exc:
            raise ValueError(f"Invalid or unreadable PPTX: {exc}") from exc

    @classmethod
    def _add_shape(
        cls,
        builder: DocTreeBuilder,
        shape: Any,
        document: Document,
        page: int,
        title_shape: Any,
        source_parts: list[str],
    ) -> None:
        if shape == title_shape:
            return
        location = _location(document, page, f"slide:{page}:shape:{shape.shape_id}")
        if shape.has_table:
            text = _table_text(
                [[cell.text for cell in row.cells] for row in shape.table.rows]
            )
            if text:
                builder.add_block("table", text, locations=(location,))
                source_parts.append(text)
        elif shape.has_text_frame:
            text = (shape.text or "").strip()
            if text:
                builder.add_block("paragraph", text, locations=(location,))
                source_parts.append(text)
        if hasattr(shape, "shapes"):
            for child in shape.shapes:
                cls._add_shape(
                    builder, child, document, page, title_shape, source_parts
                )


class OdpParser:
    """Read ODP pages, text objects, and speaker notes from its XML package."""

    block_kinds = frozenset({"paragraph", "table"})

    def parse(
        self,
        document: Document | bytes,
        *,
        document_title: str | None = None,
        metadata_overrides: dict[str, object] | None = None,
        source_path: str | Path | None = None,
    ) -> DocTree:
        document = _document(
            document, "odp", document_title, metadata_overrides, source_path
        )
        if not isinstance(document.source, bytes | bytearray):
            raise TypeError("OdpParser.parse expects Document[bytes]")
        try:
            with ZipFile(BytesIO(bytes(document.source))) as archive:
                if archive.getinfo("content.xml").flag_bits & 1:
                    raise ValueError("Encrypted ODP content is unsupported")
                if archive.getinfo("content.xml").file_size > 64 * 1024 * 1024:
                    raise ValueError("ODP content.xml exceeds 64 MiB")
                if (
                    archive.read("mimetype")
                    != b"application/vnd.oasis.opendocument.presentation"
                ):
                    raise ValueError("Invalid ODP: unexpected package mimetype")
                content = parse_untrusted_xml(archive.read("content.xml"))
            pages = [node for node in content.iter() if node.tag.endswith("}page")]
            builder = DocTreeBuilder(
                title=document.document_title or "Anonymous",
                source_path=document.source_path,
                metadata=document.metadata_overrides,
            )
            source_parts: list[str] = []
            for number, page in enumerate(pages, 1):
                objects = [child for child in page if not child.tag.endswith("}notes")]
                title = next(
                    (
                        self._text(obj)
                        for obj in objects
                        if any(
                            key.endswith("}class") and value == "title"
                            for key, value in obj.attrib.items()
                        )
                    ),
                    "",
                )
                builder.add_section(
                    1,
                    title or f"Slide {number}",
                    location=_location(document, number, f"slide:{number}"),
                )
                source_parts.append(title or f"Slide {number}")
                for index, obj in enumerate(objects, 1):
                    text = self._text(obj)
                    if text and text != title:
                        builder.add_block(
                            "paragraph",
                            text,
                            locations=(
                                _location(
                                    document, number, f"slide:{number}:object:{index}"
                                ),
                            ),
                        )
                        source_parts.append(text)
                notes = next(
                    (node for node in page if node.tag.endswith("}notes")), None
                )
                if notes is not None and (text := self._text(notes)):
                    builder.add_block(
                        "paragraph",
                        text,
                        locations=(
                            _location(document, number, f"slide:{number}:notes"),
                        ),
                        attrs={"role": "speaker_notes"},
                    )
                    source_parts.append(text)
            return replace(builder.build(), source="\n\n".join(source_parts))
        except (BadZipFile, KeyError) as exc:
            raise ValueError(f"Invalid or unreadable ODP: {exc}") from exc

    @staticmethod
    def _text(element: Any) -> str:
        paragraphs = [
            "".join(node.itertext()).strip()
            for node in element.iter()
            if node.tag.endswith("}p") or node.tag.endswith("}h")
        ]
        return "\n".join(paragraph for paragraph in paragraphs if paragraph)


class PptParser:
    """Convert legacy PPT locally with LibreOffice, then read the PPTX result."""

    block_kinds = PptxParser.block_kinds

    def parse(
        self,
        document: Document | bytes,
        *,
        document_title: str | None = None,
        metadata_overrides: dict[str, object] | None = None,
        source_path: str | Path | None = None,
    ) -> DocTree:
        document = _document(
            document, "ppt", document_title, metadata_overrides, source_path
        )
        if not isinstance(document.source, bytes | bytearray):
            raise TypeError("PptParser.parse expects Document[bytes]")
        executable = shutil.which("libreoffice") or shutil.which("soffice")
        if executable is None:
            raise RuntimeError(
                "PPT parsing requires local LibreOffice and the 'presentations' extra"
            )
        with TemporaryDirectory() as directory:
            path = Path(directory)
            source = path / "input.ppt"
            source.write_bytes(document.source)
            try:
                subprocess.run(
                    [
                        executable,
                        "-env:UserInstallation=" + (path / "profile").as_uri(),
                        "--headless",
                        "--convert-to",
                        "pptx",
                        "--outdir",
                        directory,
                        str(source),
                    ],
                    check=True,
                    capture_output=True,
                    timeout=60,
                )
                converted = (path / "input.pptx").read_bytes()
            except (
                subprocess.CalledProcessError,
                subprocess.TimeoutExpired,
                OSError,
            ) as exc:
                raise ValueError(f"Invalid or unreadable PPT: {exc}") from exc
        return PptxParser().parse(
            Document(
                converted,
                format="pptx",
                document_title=document.document_title,
                metadata_overrides=document.metadata_overrides,
                source_path=document.source_path,
            )
        )
