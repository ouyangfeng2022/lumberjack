"""Seeded random document generators for parsers beyond Markdown and DOCX.

Every generator is a pure function of ``random.Random`` and produces three
independent oracles the runner checks against the real parser:

1. ``reference_text`` — the visible text the generator knows it emitted, so
   token/character recall is checked against generated truth, not against
   another parser.
2. ``required_elements`` — exact element-signature counts (``block:record`` x N,
   ``section:h2`` x N, ...) asserted with equality rather than ``>=``.
3. ``allowed_error_types`` — for adversarial documents (truncated payloads,
   mutated markup) the parser must either build a valid tree or reject the
   input with one of the declared exception types. Any other outcome is a
   harness failure.
"""

from __future__ import annotations

import csv
import io
import json
import random
import sqlite3
from collections.abc import Callable
from dataclasses import dataclass, replace
from datetime import datetime
from typing import Any
from xml.sax.saxutils import escape as xml_escape
from xml.sax.saxutils import quoteattr as xml_quoteattr

import yaml

from lumberjack.parser.code import NotebookParser, SourceCodeParser, SQLParser
from lumberjack.parser.html import HTMLParser
from lumberjack.parser.records import (
    DelimitedTextParser,
    JSONLinesParser,
    JSONParser,
    LogParser,
    TextParser,
    TOMLParser,
    XMLParser,
    YAMLParser,
)
from lumberjack.parser.sqlite import SQLiteParser
from lumberjack.parser.xlsx import XlsxParser

ADVERSARIAL_RATIO = 0.15

_WORDS = (
    "lumber",
    "jack",
    "parser",
    "chunk",
    "section",
    "budget",
    "token",
    "heading",
    "table",
    "record",
    "provenance",
    "splitter",
    "inline",
    "block",
    "tree",
    "中文",
    "Ελληνικά",
    "العربية",
    "naïve",
    "café",
    "👩🏽‍💻",
    "Zürich",
)

_SENTENCE_SEPARATORS = (". ", ", ", "; ", " — ", " ", " ")


@dataclass(frozen=True, slots=True)
class RandomDocument:
    """One generated document plus the oracles its parser must satisfy."""

    format: str
    document_id: str
    source: str | bytes
    reference_text: str = ""
    required_elements: tuple[str, ...] = ()
    forbidden_elements: tuple[str, ...] = ()
    min_token_recall: float = 1.0
    allowed_error_types: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class FormatSpec:
    """Parser factory plus the generator producing documents for it."""

    make_parser: Callable[[], Any]
    generate: Callable[[random.Random], RandomDocument]


class _Oracle:
    """Mutable bookkeeping shared by the generators."""

    def __init__(self, rng: random.Random, prefix: str) -> None:
        self.rng = rng
        self.prefix = prefix
        self.counter = 0
        self.visible: list[str] = []

    def sentinel(self) -> str:
        self.counter += 1
        return f"{self.prefix}{self.counter:03d}"

    def sentinels(self) -> list[str]:
        return [f"{self.prefix}{number:03d}" for number in range(1, self.counter + 1)]

    def words(self, minimum: int = 1, maximum: int = 6) -> list[str]:
        count = self.rng.randint(minimum, maximum)
        return [self.rng.choice(_WORDS) for _ in range(count)]

    def sentence(self, minimum: int = 2, maximum: int = 5) -> str:
        words = self.words(minimum, maximum)
        parts: list[str] = []
        for index, word in enumerate(words):
            parts.append(word)
            if index + 1 < len(words):
                parts.append(self.rng.choice(_SENTENCE_SEPARATORS))
        return "".join(parts).strip()


def _maybe_adversarial(
    document: RandomDocument,
    rng: random.Random,
    mutate: Callable[[str, random.Random], str],
) -> RandomDocument:
    """Optionally damage the document; oracles are then relaxed to no-crash."""
    if rng.random() >= ADVERSARIAL_RATIO:
        return document
    source = document.source
    assert isinstance(source, str)
    return replace(
        document,
        source=mutate(source, rng),
        reference_text="",
        required_elements=(),
        min_token_recall=0.0,
        allowed_error_types=("ValueError",),
    )


# ---------------------------------------------------------------------------
# HTML


_HTML_ENTITIES = {"&": "&amp;", "<": "&lt;", ">": "&gt;", '"': "&quot;"}


def _html_escape(text: str) -> str:
    return "".join(_HTML_ENTITIES.get(character, character) for character in text)


def _generate_html_inline(rng: random.Random, oracle: _Oracle) -> tuple[str, str]:
    """Render one inline run and return ``(markup, visible_text)``."""
    pieces: list[str] = []
    visible: list[str] = []
    for _ in range(rng.randint(1, 4)):
        choice = rng.random()
        text = f"{oracle.sentinel()} {oracle.sentence()}"
        if choice < 0.55:
            pieces.append(_html_escape(text))
            visible.append(text)
        elif choice < 0.65:
            pieces.append(f"<strong>{_html_escape(text)}</strong>")
            visible.append(text)
        elif choice < 0.75:
            pieces.append(f"<em>{_html_escape(text)}</em>")
            visible.append(text)
        elif choice < 0.82:
            pieces.append(f"<code>{_html_escape(text)}</code>")
            visible.append(text)
        elif choice < 0.9:
            destination = f"https://example.com/{rng.randrange(1000)}"
            pieces.append(f'<a href="{destination}">{_html_escape(text)}</a>')
            visible.append(text)
        else:
            alt = f"{oracle.sentinel()} alt"
            pieces.append(
                f'<img src="image{oracle.counter}.png" alt="{_html_escape(alt)}">'
            )
            visible.append(alt)
        if rng.random() < 0.2:
            pieces.append("<br>")
    return " ".join(pieces), " ".join(visible)


def _generate_html(rng: random.Random) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_H_")
    required: list[str] = []
    body: list[str] = []
    wrapped = rng.random() >= 0.35

    if wrapped:
        body.extend(["<!DOCTYPE html>", "<html>", "<head>"])
        if rng.random() < 0.8:
            body.append(f"<title>{_html_escape(oracle.sentence())} site</title>")
        if rng.random() < 0.6:
            body.append(
                f'<meta name="description" content="{_html_escape(oracle.sentence())}">'
            )
        if rng.random() < 0.3:
            body.append("<style>body { color: red; }</style>")
        body.extend(["</head>", "<body>"])
    elif rng.random() < 0.4:
        body.append(f"<title>{_html_escape(oracle.sentence())} site</title>")

    level = 1
    for _ in range(rng.randint(4, 24)):
        choice = rng.random()
        if choice < 0.22:
            level = min(6, max(1, level + rng.choice([-2, -1, 0, 1, 1, 2, 3])))
            markup, visible = _generate_html_inline(rng, oracle)
            body.append(f"<h{level}>{markup}</h{level}>")
            oracle.visible.append(visible)
            required.append(f"section:h{level}")
        elif choice < 0.48:
            markup, visible = _generate_html_inline(rng, oracle)
            body.append(f"<p>{markup}</p>")
            oracle.visible.append(visible)
            required.append("block:paragraph")
        elif choice < 0.58:
            lines = [
                f"{oracle.sentinel()} {oracle.sentence()}"
                for _ in range(rng.randint(1, 3))
            ]
            body.append(
                "<blockquote>\n"
                + "\n".join("  " + _html_escape(line) for line in lines)
                + "\n</blockquote>"
            )
            oracle.visible.extend(lines)
            required.append("block:blockquote")
        elif choice < 0.68:
            literal = "\n".join(
                "  ".join(oracle.words(2, 4)) for _ in range(rng.randint(1, 4))
            )
            body.append("<pre>\n" + _html_escape(literal) + "\n</pre>")
            oracle.visible.append(literal)
            required.append("block:code_block")
        elif choice < 0.8:
            tag = "ol" if rng.random() < 0.5 else "ul"
            body.append(f"<{tag}>")
            for _ in range(rng.randint(2, 6)):
                item_text = f"{oracle.sentinel()} {oracle.sentence()}"
                nested = ""
                if rng.random() < 0.3:
                    nested_items = [
                        f"{oracle.sentinel()} {oracle.sentence()}"
                        for _ in range(rng.randint(1, 3))
                    ]
                    nested_tag = "ul" if rng.random() < 0.5 else "ol"
                    nested = (
                        f"<{nested_tag}>"
                        + "".join(
                            f"<li>{_html_escape(item)}</li>" for item in nested_items
                        )
                        + f"</{nested_tag}>"
                    )
                    oracle.visible.extend(nested_items)
                body.append(f"  <li>{_html_escape(item_text)}{nested}</li>")
                oracle.visible.append(item_text)
            body.append(f"</{tag}>")
            required.append("block:list")
        elif choice < 0.9:
            body.append("<table>")
            for _ in range(rng.randint(2, 5)):
                body.append("<tr>")
                for _ in range(rng.randint(1, 4)):
                    cell = f"{oracle.sentinel()} {oracle.sentence()}"
                    body.append(f"<td>{_html_escape(cell)}</td>")
                    oracle.visible.append(cell)
                body.append("</tr>")
            body.append("</table>")
            required.append("block:html_table")
        else:
            # Definition-list and bare-text runs collapse into implicit
            # paragraphs; explicit <p> boundaries pin their exact block count.
            terms = [
                f"{oracle.sentinel()} {oracle.sentence()}"
                for _ in range(rng.randint(1, 3))
            ]
            body.append(
                "<dl>"
                + "".join(f"<dt>{_html_escape(term)}</dt>" for term in terms)
                + "</dl>"
            )
            oracle.visible.extend(terms)
            bare = f"{oracle.sentinel()} {oracle.sentence()}"
            body.append(bare)
            oracle.visible.append(bare)
            first = oracle.sentence()
            body.append(f"<p>{first}</p>")
            oracle.visible.append(first)
            second = oracle.sentence()
            body.append(f"<p>{second}</p>")
            oracle.visible.append(second)
            required.extend(["block:paragraph"] * 3)
        if rng.random() < 0.25:
            body.append("<!-- generated comment -->")
        if rng.random() < 0.15:
            body.append("<script>var hidden = 'nothing to see';</script>")

    if wrapped:
        body.append("</body>\n</html>")

    def mutate(source: str, mutator: random.Random) -> str:
        mutation = mutator.randrange(4)
        if mutation == 0 and len(source) > 20:
            return source[: mutator.randint(len(source) // 4, int(len(source) * 0.95))]
        if mutation == 1:
            return source.replace("</p>", "", 1)
        if mutation == 2:
            return source.replace("<li>", '<li class=odd data-x="1">', 1)
        return source.replace("</table>", "</td></tr></table>", 1)

    return _maybe_adversarial(
        RandomDocument(
            format="html",
            document_id="",
            source="\n".join(body) + "\n",
            reference_text="\n".join(oracle.visible),
            required_elements=tuple(required),
        ),
        rng,
        mutate,
    )


# ---------------------------------------------------------------------------
# Plain text and logs


def _generate_text(rng: random.Random, *, line_mode: bool) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_T_")
    paragraphs = [
        [f"{oracle.sentinel()} {oracle.sentence()}" for _ in range(rng.randint(1, 4))]
        for _ in range(rng.randint(1, 8))
    ]
    if line_mode:
        lines = [line for paragraph in paragraphs for line in paragraph]
        source = "\n".join(lines) + "\n"
        visible = lines
        count = len(lines)
    else:
        blocks = ["\n".join(paragraph) for paragraph in paragraphs]
        source = "\n\n\n".join(blocks) + "\n"
        visible = blocks
        count = len(blocks)
    return RandomDocument(
        format="text-lines" if line_mode else "text",
        document_id="",
        source=source,
        reference_text="\n".join(visible),
        required_elements=("block:paragraph",) * count,
    )


def _generate_log(rng: random.Random) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_L_")
    lines: list[str] = []
    for index in range(rng.randint(2, 30)):
        stamp = (
            f"2026-08-{rng.randint(1, 28):02d}T"
            f"{rng.randint(0, 23):02d}:{index % 60:02d}:00Z"
        )
        level = rng.choice(["INFO", "WARN", "ERROR", "DEBUG"])
        lines.append(
            f"{stamp} {level} service-{rng.randrange(100)} "
            f"{oracle.sentinel()} {oracle.sentence()}"
        )
        if rng.random() < 0.15:
            lines.append("")
    non_empty = [line for line in lines if line]

    def mutate(source: str, mutator: random.Random) -> str:
        cut = mutator.randint(0, max(0, len(source) - 1))
        return source[:cut]

    return _maybe_adversarial(
        RandomDocument(
            format="log",
            document_id="",
            source="\n".join(lines) + "\n",
            reference_text="\n".join(non_empty),
            required_elements=("block:record",) * len(non_empty),
        ),
        rng,
        mutate,
    )


# ---------------------------------------------------------------------------
# Delimited text


def _generate_delimited(rng: random.Random, *, delimiter: str) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_C_")
    columns = rng.randint(2, 5)
    header = [
        "" if rng.random() < 0.12 else f"col_{index}_{rng.choice(_WORDS)}"
        for index in range(columns)
    ]
    normalized_header = [
        name or f"column_{index}" for index, name in enumerate(header, start=1)
    ]
    row_count = rng.randint(1, 20)
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter=delimiter, lineterminator="\n")
    writer.writerow(header)
    for _ in range(row_count):
        row: list[Any] = []
        for _ in range(columns):
            choice = rng.random()
            if choice < 0.15:
                row.append("")
            elif choice < 0.35:
                row.append(rng.randint(-9999, 9999))
            elif choice < 0.5:
                row.append(f"{oracle.sentinel()} {oracle.sentence()}")
            elif choice < 0.6:
                row.append(f"quoted {delimiter} value {rng.randrange(100)}")
            elif choice < 0.7:
                row.append(f"line one\nline two {rng.randrange(100)}")
            else:
                row.append(oracle.sentence())
        if rng.random() < 0.12:
            row.append(f"extra {oracle.sentinel()}")
        writer.writerow(row)
    buffer.seek(0)
    reader = csv.reader(buffer, delimiter=delimiter)
    next(reader, None)
    visible = list(normalized_header)
    for record in reader:
        visible.extend(str(value) for value in record)
    return RandomDocument(
        format="tsv" if delimiter == "\t" else "csv",
        document_id="",
        source=buffer.getvalue(),
        reference_text="\n".join(visible),
        required_elements=("block:tabular_row",) * row_count,
    )


# ---------------------------------------------------------------------------
# Structured data (JSON / YAML / TOML) and JSON Lines


def _random_scalar(rng: random.Random, oracle: _Oracle, *, allow_null: bool) -> Any:
    choice = rng.random()
    if choice < 0.4:
        return f"{oracle.sentinel()} {oracle.sentence()}"
    if choice < 0.55:
        return rng.randint(-100_000, 100_000)
    if choice < 0.7:
        return round(rng.uniform(-1000, 1000), 6)
    if choice < 0.85:
        return rng.choice([True, False])
    if allow_null:
        return None
    return oracle.sentence()


def _random_structure(
    rng: random.Random,
    oracle: _Oracle,
    depth: int,
    *,
    allow_null: bool,
    allow_lists: bool = True,
) -> Any:
    if depth <= 0 or rng.random() < 0.5:
        return _random_scalar(rng, oracle, allow_null=allow_null)
    if allow_lists and rng.random() < 0.4:
        return [
            _random_structure(
                rng, oracle, depth - 1, allow_null=allow_null, allow_lists=allow_lists
            )
            for _ in range(rng.randint(1, 4))
        ]
    # Bare TOML keys are restricted to ASCII letters, digits, dashes, and
    # underscores, so key words avoid the unicode pool.
    key_words = ("alpha", "beta", "gamma", "delta", "omega", "sigma", "value", "item")
    return {
        f"key_{index}_{rng.choice(key_words)}": _random_structure(
            rng, oracle, depth - 1, allow_null=allow_null, allow_lists=allow_lists
        )
        for index in range(rng.randint(1, 4))
    }


def _structure_scalars(value: Any) -> list[Any]:
    """Independent leaf count matching the documented scalar-record contract."""
    if isinstance(value, dict):
        if not value:
            return [value]
        leaves: list[Any] = []
        for child in value.values():
            leaves.extend(_structure_scalars(child))
        return leaves
    if isinstance(value, list):
        if not value:
            return [value]
        leaves = []
        for child in value:
            leaves.extend(_structure_scalars(child))
        return leaves
    return [value]


def _generate_structured(rng: random.Random, format: str) -> RandomDocument:
    allow_null = format != "toml"
    oracle = _Oracle(rng, f"LJ_{format[:2].upper()}_")
    value: Any = _random_structure(
        rng,
        oracle,
        rng.randint(1, 3),
        allow_null=allow_null,
        allow_lists=allow_null,
    )
    if not isinstance(value, dict):
        value = {"root": value}
    if format == "json":
        source = json.dumps(value, ensure_ascii=False, indent=rng.choice([None, 2]))
    elif format == "yaml":
        source = yaml.safe_dump(value, allow_unicode=True, sort_keys=False)
    else:
        source = _render_toml(value)
    leaves = _structure_scalars(value)
    reference = "\n".join(
        json.dumps(leaf, ensure_ascii=False, sort_keys=True, default=str)
        for leaf in leaves
    )

    def mutate(payload: str, mutator: random.Random) -> str:
        cut = mutator.randint(0, max(0, len(payload) - 1))
        return payload[:cut]

    return _maybe_adversarial(
        RandomDocument(
            format=format,
            document_id="",
            source=source,
            reference_text=reference,
            required_elements=("block:record",) * len(leaves),
        ),
        rng,
        mutate,
    )


def _render_toml(value: dict[str, Any]) -> str:
    lines: list[str] = []

    def render_scalar(scalar: Any) -> str:
        if isinstance(scalar, bool):
            return "true" if scalar else "false"
        if isinstance(scalar, int | float):
            return repr(scalar)
        return json.dumps(str(scalar), ensure_ascii=False)

    def emit_table(table: dict[str, Any], prefix: list[str]) -> None:
        for key, scalar in table.items():
            if isinstance(scalar, list):
                rendered = "[" + ", ".join(render_scalar(item) for item in scalar) + "]"
                lines.append(f"{key} = {rendered}")
        for key, scalar in table.items():
            if not isinstance(scalar, (dict, list)):
                lines.append(f"{key} = {render_scalar(scalar)}")
        for key, item in table.items():
            if isinstance(item, dict):
                path = [*prefix, key]
                lines.append("")
                lines.append("[" + ".".join(path) + "]")
                emit_table(item, path)

    emit_table(value, [])
    return "\n".join(lines) + "\n"


def _generate_jsonl(rng: random.Random) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_JL_")
    lines: list[str] = []
    canonical: list[str] = []
    for _ in range(rng.randint(1, 20)):
        if rng.random() < 0.12:
            lines.append("")
            continue
        value = _random_structure(rng, oracle, rng.randint(0, 2), allow_null=True)
        lines.append(json.dumps(value, ensure_ascii=False))
        canonical.append(
            json.dumps(
                value, ensure_ascii=False, sort_keys=True, separators=(", ", ": ")
            )
        )

    def mutate(source: str, mutator: random.Random) -> str:
        return source[: mutator.randint(0, max(0, len(source) - 1))]

    return _maybe_adversarial(
        RandomDocument(
            format="jsonl",
            document_id="",
            source="\n".join(lines) + "\n",
            reference_text="\n".join(canonical),
            required_elements=("block:record",) * len(canonical),
        ),
        rng,
        mutate,
    )


# ---------------------------------------------------------------------------
# XML


_XML_TAGS = ("item", "entry", "record", "node", "value", "section", "chapter")


def _generate_xml(rng: random.Random) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_X_")
    pieces: list[str] = []

    def render_element(depth: int) -> str:
        tag = rng.choice(_XML_TAGS)
        attributes = ""
        if rng.random() < 0.4:
            attributes = f" id={xml_quoteattr(f'id-{rng.randrange(10_000)}')}"
        children = rng.randint(0, 3) if depth > 0 else 0
        if children == 0:
            text = f"{oracle.sentinel()} {oracle.sentence()}"
            pieces.append(text)
            return f"<{tag}{attributes}>{xml_escape(text)}</{tag}>"
        parts = [f"<{tag}{attributes}>"]
        if rng.random() < 0.35:
            lead = f"{oracle.sentinel()} {oracle.sentence()}"
            pieces.append(lead)
            parts.append("  " + xml_escape(lead))
        for _ in range(children):
            parts.append("  " + render_element(depth - 1))
            if rng.random() < 0.25:
                tail = f"{oracle.sentinel()} {oracle.sentence()}"
                pieces.append(tail)
                parts.append("  " + xml_escape(tail))
            else:
                parts.append("  ")
        parts.append(f"</{tag}>")
        return "\n".join(parts)

    root_tag = rng.choice(_XML_TAGS)
    body = "\n".join(
        "  " + render_element(rng.randint(1, 3)) for _ in range(rng.randint(1, 5))
    )
    if rng.random() < 0.5:
        body = "<!-- document level comment -->\n" + body
    source = (
        f'<?xml version="1.0" encoding="UTF-8"?>\n<{root_tag}>\n{body}\n</{root_tag}>\n'
    )

    def mutate(payload: str, mutator: random.Random) -> str:
        return payload[: mutator.randint(0, max(0, len(payload) - 1))]

    return _maybe_adversarial(
        RandomDocument(
            format="xml",
            document_id="",
            source=source,
            reference_text="\n".join(pieces),
            required_elements=("block:record",) * len(pieces),
        ),
        rng,
        mutate,
    )


# ---------------------------------------------------------------------------
# Spreadsheets and SQLite


def _generate_xlsx(rng: random.Random) -> RandomDocument:
    from openpyxl import Workbook

    oracle = _Oracle(rng, "LJ_XL_")
    workbook = Workbook()
    visible: list[str] = []
    row_total = 0
    for sheet_index in range(rng.randint(1, 3)):
        worksheet = workbook.active if sheet_index == 0 else workbook.create_sheet()
        worksheet.title = f"Sheet{sheet_index + 1}"
        columns = rng.randint(2, 5)
        header = [
            "" if rng.random() < 0.12 else f"head_{index}" for index in range(columns)
        ]
        worksheet.append(header)
        header_names = [
            name or f"column_{index}" for index, name in enumerate(header, start=1)
        ]
        sheet_visible: list[str] = []
        sheet_rows = 0
        for _ in range(rng.randint(1, 12)):
            if rng.random() < 0.1:
                worksheet.append([])
                continue
            row: list[Any] = []
            for _ in range(columns):
                choice = rng.random()
                if choice < 0.15:
                    row.append(None)
                elif choice < 0.35:
                    row.append(rng.randint(-9999, 9999))
                elif choice < 0.5:
                    row.append(round(rng.uniform(-999, 999), 3))
                elif choice < 0.6:
                    row.append(
                        datetime(
                            2026,
                            rng.randint(1, 12),
                            rng.randint(1, 28),
                            rng.randint(0, 23),
                            rng.randint(0, 59),
                        )
                    )
                elif choice < 0.7:
                    row.append(rng.choice([True, False]))
                else:
                    row.append(f"{oracle.sentinel()} {oracle.sentence()}")
            if rng.random() < 0.1:
                row.append(f"extra {oracle.sentinel()}")
            if all(value is None for value in row):
                row[0] = f"{oracle.sentinel()} {oracle.sentence()}"
            worksheet.append(row)
            sheet_rows += 1
            for value in row:
                if value is None:
                    continue
                sheet_visible.append(
                    value.isoformat() if isinstance(value, datetime) else str(value)
                )
        if sheet_rows:
            # The parser skips sheets that contain only a header row.
            visible.extend(header_names)
            visible.extend(sheet_visible)
            row_total += sheet_rows
    stream = io.BytesIO()
    workbook.save(stream)
    source = stream.getvalue()

    def mutate(payload: bytes, mutator: random.Random) -> bytes:
        if len(payload) < 2:
            return payload
        return payload[: mutator.randint(1, len(payload) - 1)]

    return _maybe_adversarial_bytes(
        RandomDocument(
            format="xlsx",
            document_id="",
            source=source,
            reference_text="\n".join(visible),
            required_elements=("block:tabular_row",) * row_total,
        ),
        rng,
        mutate,
        ("BadZipFile", "ValueError", "KeyError"),
    )


def _generate_sqlite(rng: random.Random) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_DB_")
    connection = sqlite3.connect(":memory:")
    visible: list[str] = []
    row_total = 0
    try:
        for table_index in range(rng.randint(1, 3)):
            table = f"table_{table_index + 1}"
            columns = rng.randint(2, 5)
            names = [f"col_{index}" for index in range(columns)]
            kinds = [rng.choice(["TEXT", "INTEGER", "REAL"]) for _ in range(columns)]
            definition = ", ".join(
                f'"{name}" {kind}' for name, kind in zip(names, kinds, strict=True)
            )
            connection.execute(f'CREATE TABLE "{table}" ({definition})')
            visible.extend(names)
            for _ in range(rng.randint(1, 12)):
                values: list[Any] = []
                for kind in kinds:
                    choice = rng.random()
                    if kind == "INTEGER":
                        values.append(
                            rng.randint(-99_999, 99_999) if choice < 0.85 else None
                        )
                    elif kind == "REAL":
                        values.append(
                            round(rng.uniform(-9999, 9999), 4)
                            if choice < 0.85
                            else None
                        )
                    else:
                        values.append(
                            f"{oracle.sentinel()} {oracle.sentence()}"
                            if choice < 0.8
                            else None
                        )
                connection.execute(
                    f'INSERT INTO "{table}" VALUES ({", ".join("?" * len(values))})',
                    values,
                )
                row_total += 1
                visible.extend(str(value) for value in values if value is not None)
        connection.commit()
        source = connection.serialize()
    finally:
        connection.close()

    def mutate(payload: bytes, mutator: random.Random) -> bytes:
        if len(payload) < 2:
            return payload
        return payload[: mutator.randint(1, len(payload) - 1)]

    return _maybe_adversarial_bytes(
        RandomDocument(
            format="sqlite",
            document_id="",
            source=bytes(source),
            reference_text="\n".join(visible),
            required_elements=("block:tabular_row",) * row_total,
        ),
        rng,
        mutate,
        ("DatabaseError", "ValueError"),
    )


def _maybe_adversarial_bytes(
    document: RandomDocument,
    rng: random.Random,
    mutate: Callable[[bytes, random.Random], bytes],
    allowed_error_types: tuple[str, ...],
) -> RandomDocument:
    if rng.random() >= ADVERSARIAL_RATIO:
        return document
    source = document.source
    assert isinstance(source, bytes)
    return replace(
        document,
        source=mutate(source, rng),
        reference_text="",
        required_elements=(),
        min_token_recall=0.0,
        allowed_error_types=allowed_error_types,
    )


# ---------------------------------------------------------------------------
# Source code, notebooks, SQL


def _generate_python(rng: random.Random) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_PY_")
    lines = ['"""Module docstring outside any symbol."""', "import os", ""]
    symbols = 0
    for index in range(rng.randint(2, 12)):
        choice = rng.random()
        if choice < 0.6:
            lines.append(f"def function_{index}(first, second=1):")
            lines.append(f"    total = {oracle.sentinel()!r} + {rng.randrange(100)}")
            lines.append('    return f"{first} {second} {total}"')
            lines.append("")
            symbols += 1
        else:
            lines.append(f"class Class_{index}:")
            lines.append(f"    attribute = {oracle.sentinel()!r}")
            lines.append("")
            lines.append("    def method(self):")
            lines.append(f"        return {oracle.sentinel()!r}")
            lines.append("")
            symbols += 1
    return RandomDocument(
        format="python",
        document_id="",
        source="\n".join(lines),
        reference_text="\n".join(oracle.sentinels()),
        required_elements=("block:record",) * symbols,
    )


def _generate_javascript(
    rng: random.Random, *, typescript: bool = False
) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_TS_" if typescript else "LJ_JS_")
    lines = ["// module comment", ""]
    symbols = 0
    for index in range(rng.randint(2, 12)):
        choice = rng.random()
        if choice < 0.5:
            lines.append(f"function fn{index}(first, second) {{")
            lines.append(f"  const marker = {oracle.sentinel()!r};")
            lines.append("  return first + second + marker;")
            lines.append("}")
            lines.append("")
            symbols += 1
        elif choice < 0.75:
            lines.append(f"class Klass{index} {{")
            lines.append(f"  attribute = {oracle.sentinel()!r};")
            lines.append("")
            lines.append("  method() {")
            lines.append(f"    return {oracle.sentinel()!r};")
            lines.append("  }")
            lines.append("}")
            lines.append("")
            symbols += 1
        elif choice < 0.9 or not typescript:
            keyword = rng.choice(["const", "let", "var"])
            lines.append(f"{keyword} value{index} = {oracle.sentinel()!r};")
            lines.append("")
            symbols += 1
        else:
            lines.append(f"interface Shape{index} {{")
            lines.append(f"  label = {oracle.sentinel()!r};")
            lines.append("}")
            lines.append("")
            symbols += 1
    return RandomDocument(
        format="typescript" if typescript else "javascript",
        document_id="",
        source="\n".join(lines),
        reference_text="\n".join(oracle.sentinels()),
        required_elements=("block:record",) * symbols,
    )


def _generate_notebook(rng: random.Random) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_NB_")
    cells: list[dict[str, Any]] = []
    visible: list[str] = []
    non_empty = 0
    for _ in range(rng.randint(2, 10)):
        if rng.random() < 0.1:
            cells.append({"cell_type": "raw", "metadata": {}, "source": ""})
            continue
        body = [f"# {oracle.sentinel()}", oracle.sentence()]
        if rng.random() < 0.5:
            # nbformat line lists carry their trailing newline per element.
            source: Any = [part + "\n" for part in body]
            text = "".join(source)
        else:
            source = "\n".join(body)
            text = source
        visible.append(text)
        non_empty += 1
        cells.append(
            {
                "cell_type": rng.choice(["code", "markdown", "raw"]),
                "metadata": {},
                "source": source,
            }
        )
    notebook = {
        "cells": cells,
        "metadata": {"kernelspec": {"language": "python"}},
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    return RandomDocument(
        format="notebook",
        document_id="",
        source=json.dumps(notebook, ensure_ascii=False, indent=1),
        reference_text="\n".join(visible),
        required_elements=("block:record",) * non_empty,
    )


def _generate_sql(rng: random.Random) -> RandomDocument:
    oracle = _Oracle(rng, "LJ_SQL_")
    statements: list[str] = []
    for index in range(rng.randint(2, 12)):
        choice = rng.random()
        if choice < 0.3:
            statement = (
                f"CREATE TABLE t{index} (id INTEGER PRIMARY KEY, name TEXT, score REAL)"
            )
        elif choice < 0.6:
            literal = f"value {oracle.sentinel()} with ; semicolon and '' quote"
            statement = f"INSERT INTO t{index % 3} VALUES ({index}, '{literal}', 0.5)"
        elif choice < 0.85:
            statement = (
                f"SELECT id, name FROM t{index % 3} "
                f"WHERE note = 'a;b' AND score > {rng.random():.3f}"
            )
        else:
            statement = (
                f"UPDATE t{index % 3} SET name = $$dollar; quoted$$ WHERE id = {index}"
            )
        statements.append(statement)
    chunks: list[str] = []
    for position, statement in enumerate(statements):
        if rng.random() < 0.3:
            chunks.append(f"-- comment {oracle.sentinel()} for the next statement\n")
        chunks.append(statement)
        # Only the final statement may omit its terminating semicolon.
        if rng.random() < 0.9 or position + 1 < len(statements):
            chunks.append(";")
        chunks.append("\n" if rng.random() < 0.6 else "\n\n")
    return RandomDocument(
        format="sql",
        document_id="",
        source="".join(chunks),
        reference_text="\n".join(statements),
        required_elements=("block:record",) * len(statements),
    )


# ---------------------------------------------------------------------------
# Registry


FORMAT_SPECS: dict[str, FormatSpec] = {
    "html": FormatSpec(HTMLParser, _generate_html),
    "text": FormatSpec(TextParser, lambda rng: _generate_text(rng, line_mode=False)),
    "text-lines": FormatSpec(
        lambda: TextParser(boundary="line"),
        lambda rng: _generate_text(rng, line_mode=True),
    ),
    "log": FormatSpec(LogParser, _generate_log),
    "csv": FormatSpec(
        DelimitedTextParser, lambda rng: _generate_delimited(rng, delimiter=",")
    ),
    "tsv": FormatSpec(
        lambda: DelimitedTextParser(delimiter="\t"),
        lambda rng: _generate_delimited(rng, delimiter="\t"),
    ),
    "jsonl": FormatSpec(JSONLinesParser, _generate_jsonl),
    "json": FormatSpec(JSONParser, lambda rng: _generate_structured(rng, "json")),
    "yaml": FormatSpec(YAMLParser, lambda rng: _generate_structured(rng, "yaml")),
    "toml": FormatSpec(TOMLParser, lambda rng: _generate_structured(rng, "toml")),
    "xml": FormatSpec(XMLParser, _generate_xml),
    "xlsx": FormatSpec(XlsxParser, _generate_xlsx),
    "sqlite": FormatSpec(SQLiteParser, _generate_sqlite),
    "python": FormatSpec(lambda: SourceCodeParser(language="python"), _generate_python),
    "javascript": FormatSpec(
        lambda: SourceCodeParser(language="javascript"),
        lambda rng: _generate_javascript(rng),
    ),
    "typescript": FormatSpec(
        lambda: SourceCodeParser(language="typescript"),
        lambda rng: _generate_javascript(rng, typescript=True),
    ),
    "notebook": FormatSpec(NotebookParser, _generate_notebook),
    "sql": FormatSpec(SQLParser, _generate_sql),
}


def generate_documents(
    formats: list[str],
    *,
    seed: int,
    count_per_format: int,
) -> list[RandomDocument]:
    """Deterministically generate documents for every requested format."""
    unknown = sorted(name for name in formats if name not in FORMAT_SPECS)
    if unknown:
        raise ValueError(f"unknown random corpus formats: {unknown}")
    documents: list[RandomDocument] = []
    for name in formats:
        spec = FORMAT_SPECS[name]
        for index in range(count_per_format):
            rng = random.Random(f"{seed}:{name}:{index}".encode())
            document = spec.generate(rng)
            documents.append(replace(document, document_id=f"{name}-{index:06d}"))
    return documents


__all__ = [
    "ADVERSARIAL_RATIO",
    "FORMAT_SPECS",
    "FormatSpec",
    "RandomDocument",
    "generate_documents",
]
